import openpyxl
from typing import Dict, Any, List
from app.ingestion.parsers.base import BaseParser, ParsingResult, ParsedSourceRecord

class ExcelParser(BaseParser):
    def infer_record_type(self, sheet_name: str, headers: List[str]) -> str:
        s_lower = sheet_name.strip().lower()
        if "policy" in s_lower:
            return "POLICY"
        if "provider" in s_lower:
            return "PROVIDER"
        if "item" in s_lower or "line" in s_lower:
            return "INVOICE_LINE_ITEM"
        if "invoice" in s_lower or "bill" in s_lower:
            return "INVOICE"
        if "settle" in s_lower:
            return "SETTLEMENT"
        if "history" in s_lower:
            return "CLAIM_HISTORY"
        if "claim" in s_lower:
            return "CLAIM"
        
        # Fallback to headers
        h_lower = [str(h).strip().lower() for h in headers if h]
        if any("policy" in h for h in h_lower):
            return "POLICY"
        if any("provider" in h for h in h_lower):
            return "PROVIDER"
        if any(h in ["line_item_id", "service_code", "unit_price"] for h in h_lower):
            return "INVOICE_LINE_ITEM"
        if any(h in ["invoice_id", "invoice_reference"] for h in h_lower):
            return "INVOICE"
        if any(h in ["settlement_id", "approved_amount"] for h in h_lower):
            return "SETTLEMENT"
        return "CLAIM"

    def parse(self, file_path: str, filename: str, **kwargs) -> ParsingResult:
        records: List[ParsedSourceRecord] = []
        parsing_errors: List[Dict[str, Any]] = []
        detected_types = set()
        
        sheet_mapping = kwargs.get("sheet_mapping", {})  # e.g. {"Claims": "CLAIM", "Invoices": "INVOICE"}

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
        except Exception as e:
            return ParsingResult(
                success=False,
                total_rows_read=0,
                records=[],
                parsing_errors=[{
                    "source_file": filename,
                    "error": f"Failed to open Excel workbook: {str(e)}"
                }],
                detected_record_types=[]
            )

        total_rows = 0
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue

            # First non-empty row as header
            header_row_idx = 0
            headers = None
            for idx, r in enumerate(rows):
                if any(cell is not None and str(cell).strip() != "" for cell in r):
                    headers = [str(cell).strip() if cell is not None else f"col_{i}" for i, cell in enumerate(r)]
                    header_row_idx = idx + 1
                    break

            if not headers:
                continue

            rec_type = sheet_mapping.get(sheet_name) or self.infer_record_type(sheet_name, headers)
            detected_types.add(rec_type)

            for idx in range(header_row_idx, len(rows)):
                row_idx = idx + 1
                row_vals = rows[idx]
                total_rows += 1

                # Check blank row
                if not row_vals or all(cell is None or str(cell).strip() == "" for cell in row_vals):
                    records.append(ParsedSourceRecord(
                        source_file=filename,
                        source_sheet=sheet_name,
                        source_row=row_idx,
                        record_type_hint=rec_type,
                        raw_payload={},
                        is_skipped=True,
                        skip_reason=f"Blank row in sheet '{sheet_name}'"
                    ))
                    continue

                payload = {}
                for col_idx, h in enumerate(headers):
                    val = row_vals[col_idx] if col_idx < len(row_vals) else None
                    if val is not None:
                        # Convert to string to preserve leading zeros
                        val_str = str(val).strip()
                        # If openpyxl loaded a float that is actually an integer like 1001.0, clean up
                        if isinstance(val, float) and val.is_integer():
                            val_str = str(int(val))
                        payload[h] = val_str
                    else:
                        payload[h] = ""

                source_rec_id = None
                for key in ["claim_id", "policy_number", "invoice_id", "line_item_id", "settlement_id", "provider_id", "history_id", "id"]:
                    if key in payload and payload[key]:
                        source_rec_id = str(payload[key])
                        break

                records.append(ParsedSourceRecord(
                    source_file=filename,
                    source_sheet=sheet_name,
                    source_row=row_idx,
                    source_record_id=source_rec_id,
                    record_type_hint=rec_type,
                    raw_payload=payload,
                    is_skipped=False
                ))

        return ParsingResult(
            success=True,
            total_rows_read=total_rows,
            records=records,
            parsing_errors=parsing_errors,
            detected_record_types=list(detected_types)
        )
